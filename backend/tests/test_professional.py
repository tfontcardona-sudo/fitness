"""PROFESSIONAL (Centre Salut & Fitness): el otro negocio, no DQR repintado.

El dueño lo dijo con todas las letras: «la estructura y las preguntas no tienen
que parecerse en nada a DQR, pero el sistema interno es el mismo». Estos tests
blindan las dos mitades de esa frase:

· lo que CAMBIA — el modelo de negocio (Pack Premium mensual, el gimnasio como
  servicio del centro), el cuestionario, el documento del cliente y la cita de
  revisión, que en un centro con sala es una VISITA y no una videollamada;
· lo que NO PUEDE CAMBIAR — los números. Salen del mismo motor, con los mismos
  guardarraíles y los mismos filtros de alergias. Un documento distinto es una
  piel distinta, jamás una matemática distinta.
"""
import uuid

import pytest


def _db_available() -> bool:
    try:
        from sqlalchemy import create_engine, text

        from app.config import settings

        create_engine(settings.database_url).connect().execute(text("SELECT 1"))
        return True
    except Exception:
        return False


pytestmark = pytest.mark.skipif(not _db_available(), reason="Requiere PostgreSQL")


@pytest.fixture()
def db():
    from app.db import SessionLocal

    s = SessionLocal()
    yield s
    s.close()


def _marca(db, slug: str):
    from sqlalchemy import select

    from app.models import BrandConfig

    return db.scalar(select(BrandConfig).where(BrandConfig.slug == slug))


def _cliente(db, marca, nombre="Cliente del centro"):
    from app.models import Client
    from app.security import new_portal_token

    c = Client(full_name=nombre, email=f"pf-{uuid.uuid4().hex[:8]}@test.local",
               portal_token="tmp", status="active", package_tier="full",
               brand_id=marca.id)
    db.add(c)
    db.flush()
    c.portal_token = new_portal_token(c.id)
    db.commit()
    return c


# --------------------------------------------------------- el negocio ------

def test_el_pack_premium_es_mensual_y_el_gimnasio_no_es_asesoria(db):
    """Professional no vende programas cerrados de 3 o 6 meses como DQR: es
    cuota mensual con renovación. Y el plan de gimnasio (60 €) NO es asesoría
    — se enseña con lo que se cobra en el centro, no como producto de la web."""
    from app.services import branding

    pf = branding.marca_por_id(_marca(db, "professional-fitness").id, db)
    assert pf.importe("full", "1m") == 12990          # 129,90 €
    assert pf.vende() == [("full", "1m")]             # una cosa, una duración
    assert not pf.vende_oferta()
    titulos = " · ".join(s.get("title", "") for s in (pf.extra_services or []))
    assert "gimnasio" in titulos.lower()
    assert "60 €/mes" in " ".join(s.get("price", "") for s in pf.extra_services)


def test_la_renovacion_de_una_cuota_mensual_avisa_a_los_25_dias(db):
    """El ciclo que pidieron: pago mensual → seguimiento → revisión →
    renovación. Sin suscripción que se cobre sola, el aviso lo da el sistema."""
    from datetime import date, datetime, timedelta, timezone

    from app.services import renewals

    class _C:
        payment_status = "paid"
        stripe_subscription_id = None
        billing_period = "1m"
        paid_at = datetime.now(timezone.utc) - timedelta(days=26)

    ventana = renewals.renewal_window(_C(), date.today())
    assert ventana is not None
    _fin, quedan = ventana
    assert quedan <= renewals.RENEWAL_WARN_DAYS
    assert renewals.is_due(_C(), date.today())


# ----------------------------------------------------- el cuestionario -----

def test_el_cuestionario_de_professional_es_suyo_y_no_el_de_nadie_mas():
    """Lo BÁSICO y nada más: de lo que salen los números y la seguridad, más
    las tres cosas que un centro con local necesita por escrito. Y sus
    preguntas NO son las de la variante 'simple': compartirlas habría hecho
    que retocar un centro cambiara el cuestionario del otro."""
    from app.services.anamnesis_variant import definicion, etiquetas

    pf = definicion("professional")
    simple = definicion("simple")
    claves_pf = [q["key"] for q in pf["extra_questions"]]
    assert claves_pf == ["experiencia", "disponibilidad", "obstaculo"]
    assert claves_pf != [q["key"] for q in simple["extra_questions"]]
    # Un extra sí: la zona a priorizar. El resto se pregunta en la sala.
    assert pf["optional_blocks"] == ["priority_zones"]
    assert "experiencia" in etiquetas("professional")


def test_las_respuestas_del_centro_se_etiquetan_en_las_notas():
    from app.services.anamnesis_variant import anexar_respuestas

    out = anexar_respuestas("Nota previa del coach", "professional",
                            {"disponibilidad": "Lunes y miércoles a las 18 h",
                             "inventada": "esto no se pregunta"})
    assert "Nota previa del coach" in out
    assert "Lunes y miércoles a las 18 h" in out
    # Lo que la variante NO pregunta no puede escribir en la ficha: el cuerpo
    # llega del navegador.
    assert "esto no se pregunta" not in out


# -------------------------------------------------------- el documento -----

def _plan_de_prueba():
    nutricion = {
        "tdee_kcal": 2600, "target_kcal": 2200,
        "rationale": "Déficit moderado para bajar grasa sin perder fuerza.",
        "macros": {"protein_g": 165, "carbs_g": 220, "fat_g": 73},
        "meals": [{"slot": 1, "name": "Desayuno", "time": "08:00",
                   "target": {"kcal": 600, "protein_g": 45, "carbs_g": 60, "fat_g": 20}}],
        "meal_bank": {"mode": "flexible_7", "slots": [
            {"slot": 1, "fmt": "options", "options": [
                {"title": "Avena con yogur", "prep": "Mezclar.", "prep_minutes": 5,
                 "ingredients": [{"food": "avena", "grams": 80, "household": "8 cucharadas"}],
                 "macros": {"kcal": 600, "protein_g": 45, "carbs_g": 60, "fat_g": 20}}]}]},
        "supplements": [{"name": "Creatina", "dose": "5 g", "timing": "cualquier momento",
                         "evidence_note": "Evidencia sólida en fuerza"}],
        "flexibility_rules": ["Puedes cambiar el arroz por pasta al mismo peso."],
    }
    entreno = {
        "split_name": "Full Body", "split_rationale": "Tres días te cunden más repartidos.",
        "weekly_progression": [
            {"week": w, "intent": "Base", "load_pct": 100, "rir_target": "2",
             "volume_note": "Carga estable"} for w in (1, 2, 3, 4)],
        "sessions": [{"day": "Lunes", "name": "Full A",
                      "warmup": "5 min bici · movilidad",
                      "cooldown": "Estiramiento suave",
                      "exercises": [{"exercise_id": 1, "sets": 4, "rep_range": "8-10",
                                     "rir": "2", "rest_sec": 120,
                                     "progression_rule": "Completas 4×10 → +2,5 kg",
                                     "technique_cue": "Codos a 45°",
                                     "biomech_cue": "Protege el hombro"}]}],
        "cardio": {"daily_steps": 9000, "sessions": []},
        "deload_instructions": "Semana 4: series ÷2 · carga −30 %",
    }
    return nutricion, entreno


def test_el_documento_de_professional_no_se_parece_al_de_dqr():
    """Mismo plan, dos documentos distintos: el de Professional trae su portada
    y la sección del ciclo del centro, y deja fuera el material de consulta de
    DQR (índice, plato saludable, preguntas frecuentes)."""
    from io import BytesIO

    from docx import Document

    from app.services.docs.plan_doc import generate_plan_doc
    from app.services.docs.plan_doc_pf import generate_plan_doc_pf
    from app.services.docs.word_base import DocBrand

    nutricion, entreno = _plan_de_prueba()
    marca = DocBrand(name="Professional", color_primary="#C9A227",
                     color_secondary="#161616", font_family="Inter",
                     contact_address="Carretera Pierre Vilar, 2 · 17002 Girona")
    comun = dict(brand=marca, client_name="Marta", month_index=2,
                 goal_type="fat_loss", diet_mode="flexible_7",
                 nutrition=nutricion, training=entreno,
                 exercise_names={1: "Press banca"},
                 include_nutrition=True, include_training=True)

    pf = generate_plan_doc_pf(**comun)
    dqr = generate_plan_doc(**comun, education={})

    def texto(b: bytes) -> str:
        doc = Document(BytesIO(b))
        partes = [p.text for p in doc.paragraphs]
        for t in doc.tables:
            for fila in t.rows:
                partes += [c.text for c in fila.cells]
        return "\n".join(partes).upper()

    t_pf, t_dqr = texto(pf), texto(dqr)
    # Lo que SOLO tiene Professional: su portada y el ciclo del centro.
    for propio in ("TU MES EN UNA PÁGINA", "CÓMO TRABAJAMOS ESTO",
                   "DÓNDE ESTAMOS Y QUÉ TOCA AHORA", "CÓMO SE LEE TU RUTINA"):
        assert propio in t_pf, propio
        assert propio not in t_dqr, propio
    # Lo de DQR que NO recibe el cliente del centro.
    for ajeno in ("EN ESTE DOCUMENTO", "EL PLATO SALUDABLE", "ALIMENTOS POR GRUPOS"):
        assert ajeno not in t_pf, ajeno
    # Y la dirección del centro, que en un gimnasio es lo primero que se busca.
    assert "GIRONA" in t_pf


def test_las_cifras_del_documento_de_professional_son_las_mismas():
    """La piel cambia; la matemática NO. Las kcal, los macros y los gramos del
    plan son los que calculó el backend, no otros."""
    from io import BytesIO

    from docx import Document

    from app.services.docs.plan_doc_pf import generate_plan_doc_pf
    from app.services.docs.word_base import DocBrand

    nutricion, entreno = _plan_de_prueba()
    data = generate_plan_doc_pf(
        brand=DocBrand(name="Professional", color_primary="#C9A227",
                       color_secondary="#161616", font_family="Inter"),
        client_name="Marta", month_index=1, goal_type="fat_loss",
        diet_mode="flexible_7", nutrition=nutricion, training=entreno,
        exercise_names={1: "Press banca"},
        include_nutrition=True, include_training=True)
    doc = Document(BytesIO(data))
    texto = "\n".join([p.text for p in doc.paragraphs]
                      + [c.text for t in doc.tables for f in t.rows for c in f.cells])
    assert "2.200" in texto                    # las kcal, en español
    assert "165 g" in texto and "220 g" in texto and "73 g" in texto
    assert "Déficit de 400 kcal" in texto      # 2600 → 2200, el ajuste real
    assert "Press banca" in texto and "4×8-10" in texto
    assert "avena 80 g" in texto.lower()


def test_el_alergeno_tampoco_se_cuela_en_el_documento_del_centro():
    """El filtro de seguridad es el MISMO. Si una regla de flexibilidad nombra
    un alimento prohibido, no se imprime — da igual la marca."""
    from io import BytesIO

    from docx import Document

    from app.services.docs.plan_doc_pf import generate_plan_doc_pf
    from app.services.docs.word_base import DocBrand

    nutricion, entreno = _plan_de_prueba()
    nutricion["flexibility_rules"] = ["Puedes cambiar el arroz por pan integral."]
    data = generate_plan_doc_pf(
        brand=DocBrand(name="Professional", color_primary="#C9A227",
                       color_secondary="#161616", font_family="Inter"),
        client_name="Marta", month_index=1, goal_type="fat_loss",
        diet_mode="flexible_7", nutrition=nutricion, training=entreno,
        food_allergies=["gluten"], include_nutrition=True, include_training=False)
    doc = Document(BytesIO(data))
    texto = "\n".join([p.text for p in doc.paragraphs]
                      + [c.text for t in doc.tables for f in t.rows for c in f.cells])
    assert "pan integral" not in texto.lower()


def test_el_word_de_professional_se_puede_subir_editado(db):
    """«Subir Word editado» tiene que seguir funcionando con el documento del
    centro: sus barras se llaman distinto y se traducen (`_ALIAS_BARRA`). Sin
    esto el coach editaba el Word, lo subía y sus cambios se perdían MUDOS."""
    from io import BytesIO

    from docx import Document

    from app.models import Plan
    from app.services.docs.plan_doc_pf import generate_plan_doc_pf
    from app.services.docs.word_base import DocBrand
    from app.services.word_import import parse_word_edits

    nutricion, entreno = _plan_de_prueba()
    data = generate_plan_doc_pf(
        brand=DocBrand(name="Professional", color_primary="#C9A227",
                       color_secondary="#161616", font_family="Inter"),
        client_name="Marta", month_index=1, goal_type="fat_loss",
        diet_mode="flexible_7", nutrition=nutricion, training=entreno,
        exercise_names={1: "Press banca"},
        include_nutrition=True, include_training=True)

    # Edita a mano la caja de SUPLEMENTOS (que en Professional se llama así, no
    # «Suplementación recomendada») y la de la semana suave (el «deload»).
    doc = Document(BytesIO(data))
    tocadas = 0
    for t in doc.tables:
        if len(t.rows) == 1 and len(t.rows[0].cells) == 1:
            celda = t.rows[0].cells[0]
            txt = celda.text
            if txt.startswith("Creatina"):
                celda.paragraphs[0].runs[0].text = "Creatina — 5 g (después de entrenar)"
                tocadas += 1
            elif "series ÷2" in txt:
                celda.paragraphs[0].runs[0].text = "Semana 4: mitad de series"
                tocadas += 1
    assert tocadas == 2, "el documento ya no imprime esas cajas: el test no prueba nada"
    buf = BytesIO()
    doc.save(buf)

    plan = Plan(client_id=None, month_index=1, nutrition_json=nutricion,
                training_json=entreno, education_json=None)
    res = parse_word_edits(db, plan, buf.getvalue())
    pistas = " · ".join((res.get("extra_changes") or []) + (res.get("warnings") or []))
    supl = res["nutrition_json"]["supplements"][0]
    assert supl["timing"] == "después de entrenar", pistas
    assert res["training_json"]["deload_instructions"] == "Semana 4: mitad de series", pistas


# ------------------------------------------------------------ la cita ------

def test_la_revision_de_un_centro_es_una_visita_no_una_videollamada(db):
    from app.services import branding, citas

    branding.invalidar()
    pf = _marca(db, "professional-fitness")
    dqr = _marca(db, "dqr")
    assert citas.modo_de_marca(db, _cliente(db, pf)) == citas.PRESENCIAL
    assert citas.modo_de_marca(db, _cliente(db, dqr)) == citas.VIDEOLLAMADA
    assert "centro" in citas.etiqueta(citas.PRESENCIAL)
    assert citas.etiqueta(citas.VIDEOLLAMADA) == "videollamada"


def test_una_cita_ya_acordada_no_cambia_de_tipo_con_el_switch():
    """Mismo criterio que la marca del cliente: lo SELLADO en la fila manda.
    Si el centro pasara mañana a hacerlas por vídeo, la visita ya confirmada
    sigue siendo una visita."""
    from app.services import citas

    class _Fila:
        modo = "presencial"

    assert citas.es_presencial(_Fila()) is True

    class _Vieja:
        modo = None                     # las filas de siempre, sin columna

    assert citas.modo_de_cita(_Vieja()) == citas.VIDEOLLAMADA
    assert citas.es_presencial(_Vieja()) is False


def test_el_coach_puede_cerrar_una_visita_sin_google_conectado(db):
    """El ciclo completo de una VISITA, que es donde estaba el nudo: un centro
    no tiene por qué tener una cuenta de Google, y exigirla dejaba la cita
    atascada para siempre en «pendiente de agendar».

    Se ejercita por la API real: el cliente propone desde su portal y el coach
    acepta desde el panel, con Google DESCONECTADO."""
    import os
    from datetime import date, datetime, timedelta

    from fastapi.testclient import TestClient

    from app.main import app
    from app.models import Period, Plan, VideoCall
    from app.security import create_access_token
    from app.services import branding, citas

    branding.invalidar()
    pf = _marca(db, "professional-fitness")
    c = _cliente(db, pf, "Cliente que viene al centro")
    # Una revisión cerrada: es lo que abre la cita.
    plan = Plan(client_id=c.id, month_index=1, version=1, status="published",
                training_json={}, nutrition_json={})
    db.add(plan)
    db.flush()
    db.add(Period(client_id=c.id, plan_id=plan.id, period_index=1, status="analyzed",
                  starts_on=date.today() - timedelta(days=20),
                  ends_on=date.today() - timedelta(days=6)))
    db.commit()

    cuando = (datetime.now() + timedelta(days=3)).replace(microsecond=0)
    with TestClient(app) as http:
        # 1) El cliente propone desde su portal: la cita nace ya como VISITA.
        r = http.post(f"/api/p/{c.portal_token}/video-call",
                      json={"start_at": cuando.isoformat()})
        assert r.status_code == 200, r.text
        assert r.json()["modo"] == "presencial"
        assert "centro" in r.json()["modo_label"]

        db.expire_all()
        vc = db.scalar(
            __import__("sqlalchemy").select(VideoCall)
            .where(VideoCall.client_id == c.id))
        assert vc.modo == "presencial"

        # 2) El coach la acepta SIN Google. Con una videollamada esto sería un
        #    409 «conecta tu cuenta»; con una visita, se confirma.
        auth = {"Authorization": f"Bearer {create_access_token(os.environ.get('ADMIN_1_USER', 'coach1'))}"}
        r = http.post(f"/api/clients/{c.id}/video-calls/{vc.id}/accept",
                      json={"duration_min": 30}, headers=auth)
        assert r.status_code == 200, r.text
        cuerpo = r.json()
        assert cuerpo["status"] == "scheduled"
        assert cuerpo["modo"] == "presencial"
        # Una visita no tiene enlace: un Meet muerto es peor que ninguno.
        assert cuerpo["meet_url"] is None

        # 3) Y el cliente ve el día y la DIRECCIÓN, no un botón de «unirme».
        est = http.get(f"/api/p/{c.portal_token}/video-call").json()
        assert est["state"] == "scheduled" and est["modo"] == "presencial"
        assert "Girona" in (est.get("lugar") or "")
        assert est["call"]["meet_url"] is None

    db.expire_all()
    assert citas.es_presencial(db.get(VideoCall, vc.id))
