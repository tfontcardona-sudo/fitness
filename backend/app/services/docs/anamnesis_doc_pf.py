"""El documento de anamnesis de PROFESSIONAL (Centre Salut & Fitness).

No es la ficha de DQR con otro color: el orden y el énfasis cambian con el
negocio, igual que ya pasa con `plan_doc_pf` frente a `plan_doc`.

· DQR es asesoría online: la ficha completa (objetivo → entreno → dieta →
  salud → experiencia) es lo único que el coach tiene del cliente, así que
  entra todo con el mismo peso.
· Professional es un centro con sala: al cliente se le ve la cara. Lo que
  hace falta por escrito es lo que hace falta ANTES de que pise el gimnasio
  — entrenamiento y salud primero (contraindicaciones antes de programar un
  ejercicio) — y la alimentación, más corta porque el cuestionario del
  centro también lo es (§ `AnamnesisProfesional.tsx`: "lo que no hace falta
  por escrito se pregunta en persona").

Los DATOS son los mismos (misma ficha, mismos campos); lo que cambia es el
orden, el tono y la piel — negro y dorado, sin semáforo de colores.
"""

from __future__ import annotations

import io
from datetime import date

from docx import Document
from docx.shared import Pt

from app.services.docs.plan_doc_pf import (
    _barra,
    _caja,
    _nota,
    _portada_generica,
    _tabla,
)
from app.services.docs.word_base import (
    DocBrand,
    add_bullets,
    init_document,
    setup_reference_pages,
    spacer,
    _hex,
)

_OBJETIVO = {
    "fat_loss": "Bajar grasa", "muscle_gain": "Ganar músculo",
    "recomp": "Bajar grasa y ganar músculo", "maintenance": "Mantener y coger el hábito",
    "injury_recovery": "Volver después de una lesión",
}
_LEVEL_LABEL = {"beginner": "Principiante", "intermediate": "Intermedio", "advanced": "Avanzado"}
_PLACE_LABEL = {"gym": "Sala del centro", "home": "Casa", "outdoor": "Exterior"}
_ACTIVITY_LABEL = {
    "sedentary": "Sedentaria (oficina)", "light": "Ligera (de pie a ratos)",
    "active": "Activa (trabajo físico)", "very_active": "Muy activa (físico intenso)",
}
_DIET_PATTERN_LABEL = {
    "vegano": "Vegano", "vegetariano": "Vegetariano", "pescetariano": "Pescetariano",
    "sin_cerdo": "Sin cerdo", "halal": "Halal", "kosher": "Kosher",
}
_SEX_LABEL = {"male": "Hombre", "female": "Mujer"}


def _f1(v: float | None, unit: str = "") -> str:
    return "—" if v is None else f"{v:.1f}{unit}".replace(".0", "").replace(".", ",")


def generate_anamnesis_doc_pf(
    *,
    brand: DocBrand,
    client_name: str,
    sex: str | None,
    age: int | None,
    height_cm: float | None,
    start_weight_kg: float | None,
    current_weight_kg: float | None,
    body_fat_pct: float | None,
    goal_type: str | None,
    goal_weight_kg: float | None,
    goal_deadline: str | None,  # noqa: ARG001 — la ficha de Professional no lo imprime (se acuerda en el mostrador)
    level: str | None,
    training_days: int | None,
    training_place: str | None,
    session_max_min: int | None,
    daily_activity_level: str | None,
    equipment: list[str] | None,
    diet_mode: str | None,  # noqa: ARG001 — Professional no distingue flexible/estricta por escrito
    diet_pattern: str | None,
    meals_per_day: int | None,  # noqa: ARG001 — se acuerda en persona
    food_allergies: list[str] | None,
    food_dislikes: list[str] | None,
    food_likes: list[str] | None,
    injuries_notes: str | None,
    medical_notes: str | None,
    medication_notes: str | None,
    current_supplements: str | None,
    sport_history: str | None,
    lifestyle_notes: str | None,
    generated_on: date | None = None,
) -> bytes:
    """Misma firma que `generate_anamnesis_doc` (DQR): quien la llama no
    prepara datos distintos, solo elige el generador."""
    doc = init_document(brand)
    hoy = generated_on or date.today()
    pie = brand.name + (f" · {brand.contact_address}" if brand.contact_address else "")
    setup_reference_pages(
        doc,
        logo_path=None,
        right_title=f"FICHA | {client_name}",
        right_sub=hoy.strftime("%d/%m/%Y"),
        footer_text=pie,
    )

    _portada_generica(
        doc, brand, client_name, rotulo="Ficha del cliente",
        datos=[("Objetivo", _OBJETIVO.get(goal_type or "", "A definir en el centro")),
               ("Nivel", _LEVEL_LABEL.get(level or "", "—")),
               ("Ficha abierta", hoy.strftime("%d/%m/%Y"))],
    )
    doc.add_page_break()

    # ------------------------------------------------- tu ficha en una página
    _barra(doc, "Su ficha en una página")
    _nota(doc, "Lo esencial para programarle, de un vistazo.")
    filas = [
        ["Sexo", _SEX_LABEL.get(sex or "", "—")],
        ["Edad", f"{age} años" if age else "—"],
        ["Altura", _f1(height_cm, " cm")],
        ["Peso", _f1(current_weight_kg if current_weight_kg is not None else start_weight_kg, " kg")],
    ]
    if body_fat_pct is not None:
        filas.append(["% graso", _f1(body_fat_pct, " %")])
    if goal_weight_kg is not None:
        filas.append(["Peso objetivo", _f1(goal_weight_kg, " kg")])
    _tabla(doc, ["Dato", "Valor"], filas, brand, anchos=[3200, 6466])

    # ----------------------------------------------------------- entrenamiento
    spacer(doc, 14)
    _barra(doc, "Entrenamiento")
    _tabla(doc, ["Dato", "Valor"], [
        ["Días/semana en el centro", str(training_days) if training_days else "—"],
        ["Dónde entrena", _PLACE_LABEL.get(training_place or "", "—")],
        ["Tiempo por sesión", f"{session_max_min} min" if session_max_min else "—"],
        ["Actividad fuera del centro", _ACTIVITY_LABEL.get(daily_activity_level or "", "—")],
        ["Material propio (fuera del centro)", ", ".join(equipment) if equipment else "—"],
    ], brand, anchos=[3800, 5866])

    # ------------------------------------------------------------------ salud
    # Antes que la dieta: lo que un entrenador necesita saber ANTES de que el
    # cliente toque una barra, no después.
    spacer(doc, 14)
    _barra(doc, "Salud")
    hay_salud = bool(injuries_notes or medical_notes or medication_notes)
    if injuries_notes:
        _caja(doc, [("⚠ Lesiones", injuries_notes)])
    if medical_notes:
        spacer(doc, 6)
        _caja(doc, [("Patologías y salud", medical_notes)])
    if medication_notes:
        spacer(doc, 6)
        _caja(doc, [("Medicación", medication_notes)])
    if not hay_salud:
        doc.add_paragraph("Sin lesiones ni patologías declaradas.")
    if current_supplements:
        spacer(doc, 6)
        doc.add_paragraph(f"Suplementación actual: {current_supplements}")

    # ------------------------------------------------------------ alimentación
    spacer(doc, 14)
    _barra(doc, "Alimentación")
    _nota(doc, "Lo que hace falta saber por escrito; el resto se acuerda en el centro.")
    p = doc.add_paragraph()
    run = p.add_run(
        "⚠ Alergias / intolerancias: " + (", ".join(food_allergies) if food_allergies else "Ninguna declarada")
    )
    run.font.bold = True
    run.font.color.rgb = _hex("A61B0E" if food_allergies else "5F5C57")
    if diet_pattern:
        doc.add_paragraph(f"Patrón dietético: {_DIET_PATTERN_LABEL.get(diet_pattern, diet_pattern)}")
    if food_dislikes:
        doc.add_paragraph("No le gustan: " + ", ".join(food_dislikes))
    if food_likes:
        doc.add_paragraph("Preferencias: " + ", ".join(food_likes))

    # -------------------------------------------------------------- experiencia
    if sport_history or lifestyle_notes:
        spacer(doc, 14)
        _barra(doc, "Experiencia y hábitos")
        if sport_history:
            doc.add_paragraph(sport_history)
        if lifestyle_notes:
            add_bullets(doc, [l for l in lifestyle_notes.splitlines() if l.strip()] or [lifestyle_notes])

    p = doc.add_paragraph()
    run = p.add_run(f"Ficha generada el {hoy.strftime('%d/%m/%Y')}")
    run.font.size = Pt(8)
    run.font.color.rgb = _hex("8A857D")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
